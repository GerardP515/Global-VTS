"""Carry the legacy v0.3 Stage 4 findings (legacy B10 to B34) onto the IMO 2025 TSS IDs.

Inputs (legacy v0.3 files are read from git history, commit 63b9000):
  research/stage4/stage4_findings_v0_3.csv, data/current/Global_VTS_Inventory_v0_3.xlsx
  research/reconciliation/TSS_Crosswalk_Legacy_v0_2_to_IMO_2025.csv
  data/authoritative/IMO_2025_Individual_TSS_Inventory.csv
Output:
  research/reconciliation/Legacy_v0_3_Findings_to_IMO_2025.csv

These are prior-research leads for Stage 2, not audited Stage 2 findings.
No VTS IDs are allocated here; that happens when the master register is built.
"""
import csv
import io
import os
import subprocess
from collections import defaultdict

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
LEGACY_COMMIT = "63b9000"
OUT = os.path.join(ROOT, "research", "reconciliation", "Legacy_v0_3_Findings_to_IMO_2025.csv")

# Legacy v0.3 scheme IDs -> IMO 2025 Part I VRS IDs (data/authoritative/IMO_2025_Mandatory_Reporting_Systems.csv)
IMO_VRS = {
    "SRS-004": ("VRS-0016", "STRAITREP"),
    "SRS-009": ("VRS-0006", "OUESSREP"),
    "SRS-010": ("VRS-0011", "FINREP"),
    "SRS-011": ("VRS-0012", "COPREP"),
    "SRS-012": ("VRS-0023", "CANREP"),
    "SRS-015": ("VRS-0013", "GIBREP"),
    "SRS-016": ("VRS-0015", "ADRIREP"),
    "SRS-022": ("VRS-0022", "Off Chengshan Jiao Promontory"),
    "SRS-025": ("VRS-0021", "Off the north-eastern and south-eastern coasts of the United States (WHALESNORTH)"),
}
# Mandatory under national law only: not in IMO Part I, so no VRS ID exists yet.
NATIONAL_MRS = {
    "SRS-017": "TUBRAP (Turkish Straits, national; IMO recommends only)",
    "SRS-020": "SUNDAREP (national; mandatory for Indonesian-flag ships only)",
    "SRS-021": "LOMBOKREP (national; mandatory for Indonesian-flag ships only)",
    "SRS-023": "MASTREP (Australia, national; foreign ships only between first arrival and final departure)",
}

# Component-level knowledge from the legacy research where one legacy row splits into several IMO 2025 TSS.
# canonical_id -> (vts_confirmed, mrs_confirmed or None to keep, note)
COMPONENT = {
    "TSS-0055": (True, None, "Legacy research: eastern 4 to 5 nm of IJmuiden West Inner lies inside the 12 nm VTS North Sea Canal Area (derived from COLREG.2/Circ.64 coordinates; reference position approximate)."),
    "TSS-0056": (False, None, "Legacy research: IJmuiden North is outside the 12 nm VTS North Sea Canal Area."),
    "TSS-0057": (False, None, "Legacy research: IJmuiden West Outer is outside the 12 nm VTS North Sea Canal Area."),
    "TSS-0173": (False, None, "Legacy research: only the south-east end of the Santa Barbara Channel TSS (near Point Vicente) is within the VTS Los Angeles-Long Beach VMRS radius; the rest is 53 to 135 nm away. Partial only, so not carried as confirmed."),
    "TSS-0174": (True, None, "Legacy research: the LA/LB approaches fall within the 25 nm VMRS radius of Point Fermin, except the tip of the southern approach (about 25.7 nm)."),
}

# Legacy records in waters where WETREP may apply; legacy research never checked it.
WETREP_CHECK = set()  # filled in main() from the France and Iberia legacy rows

# Where IMO 2025 batch research already on main (B20, 29 Sep 2026) covers the same TSS.
MAIN_B20 = {
    "TSS-0096": "Agrees with main B20 (inside COPREP; Roca Control).",
    "TSS-0097": "Agrees with main B20 (unresolved).",
    "TSS-0098": "Agrees with main B20 (GIBREP names the TSS; Tarifa and Tangier Traffic).",
    "TSS-0099": "CONFLICT with main B20: main records Almeria as a lead only and no VTS. Legacy VTS finding rests on a 2025 Ministry of Transport magazine article saying Salvamento Maritimo controls the Cabo de Gata DST. Resolve before registering.",
    "TSS-0100": "Agrees with main B20 (unresolved).",
}

TODAY = "2026-09-29"


def git_show(path, binary=False):
    out = subprocess.run(["git", "-C", ROOT, "show", f"{LEGACY_COMMIT}:{path}"], capture_output=True, check=True).stdout
    return out if binary else out.decode("utf-8")


def main():
    findings = {r["Record ID"]: r for r in csv.DictReader(io.StringIO(git_show("research/stage4/stage4_findings_v0_3.csv")))}
    wb = openpyxl.load_workbook(io.BytesIO(git_show("data/current/Global_VTS_Inventory_v0_3.xlsx", binary=True)))
    services = {r[0]: r for r in wb["VTS Services"].iter_rows(min_row=7, values_only=True) if r[0]}
    schemes = {r[0]: r for r in wb["Reporting Schemes"].iter_rows(min_row=7, values_only=True) if r[0]}
    sources = {r[0]: r for r in wb["Sources"].iter_rows(min_row=7, values_only=True) if r[0]}

    tss = {r["canonical_id"]: r for r in csv.DictReader(open(os.path.join(ROOT, "data", "authoritative", "IMO_2025_Individual_TSS_Inventory.csv"), encoding="utf-8"))}
    cw = list(csv.DictReader(open(os.path.join(ROOT, "research", "reconciliation", "TSS_Crosswalk_Legacy_v0_2_to_IMO_2025.csv"), encoding="utf-8")))
    for r in cw:
        if r["legacy_id"] in findings and r["legacy_id"] in ("TSS-0059", "TSS-0060", "TSS-0061", "TSS-0062"):
            WETREP_CHECK.add(r["canonical_id"])
    by_canon = defaultdict(list)
    legacy_split = defaultdict(int)
    for r in cw:
        legacy_split[r["legacy_id"]] += 1
    for r in cw:
        if r["legacy_id"] in findings:
            by_canon[r["canonical_id"]].append(r)

    rows = []
    for cid in sorted(by_canon):
        maps = by_canon[cid]
        recs = [findings[m["legacy_id"]] for m in maps]
        legacy_ids = "; ".join(m["legacy_id"] for m in maps)
        split = any(legacy_split[m["legacy_id"]] > 1 for m in maps)
        merged = len(maps) > 1
        superseded = any(r["Operational status"] == "Withdrawn/superseded" for r in recs)

        vts_ids = []
        for r in recs:
            if r["VTS finding"] == "Confirmed present":
                vts_ids += [x for x in r["VTS service IDs"].split("; ") if x and x not in vts_ids]
        vts_lead_ids = []
        for r in recs:
            if r["VTS finding"] != "Confirmed present":
                vts_lead_ids += [x for x in (r["VTS service IDs"] or "").split("; ") if x and x not in vts_lead_ids]
        mrs_ids = []
        for r in recs:
            if r["Reporting finding"] == "Confirmed present":
                mrs_ids += [x for x in r["Reporting scheme IDs"].split("; ") if x and x not in mrs_ids]
        vol = [x for r in recs for x in (r["Reporting scheme IDs"] or "").split("; ")
               if x and schemes.get(x) and schemes[x][3] == "Voluntary"]

        vts_ok = bool(vts_ids)
        mrs_ok = bool(mrs_ids)
        notes = []
        comp_note = None
        if cid in COMPONENT:
            v, m, comp_note = COMPONENT[cid]
            if not v:
                vts_lead_ids = vts_ids + vts_lead_ids
                vts_ids = []
            vts_ok = v and bool(vts_ids)
            if m is not None:
                mrs_ok = m
        if superseded:
            vts_ok = mrs_ok = False
            notes.append("Legacy rows described the superseded Odesa/Ilichevsk schemes; research the successor scheme directly.")

        status = {(True, True): "VTS + MRS", (True, False): "VTS only", (False, True): "MRS only"}.get((vts_ok, mrs_ok), "Unresolved")

        if split and cid not in COMPONENT:
            notes.append("Legacy finding was made at the grouped-heading level; confirm it applies to this individual TSS.")
        if merged:
            notes.append("Several legacy rows map to this TSS; findings combined.")

        imo_mrs = [IMO_VRS[x] for x in mrs_ids if x in IMO_VRS] if mrs_ok else []
        nat_mrs = [NATIONAL_MRS[x] for x in mrs_ids if x in NATIONAL_MRS] if mrs_ok else []
        region = tss[cid]["imo_section"]
        if cid in WETREP_CHECK:
            notes.append("Legacy research did not check WETREP (VRS-0005, heavy-grade oil tankers of 600 dwt and above); check it for this TSS.")
        rec_src = []
        for r in recs:
            rec_src += [x for x in r["Source IDs"].split("; ") if x and x not in rec_src]

        def svc_field(ids, i):
            return "; ".join(str(services[x][i]) for x in ids if x in services and services[x][i])

        rows.append({
            "tss_id": cid,
            "tss_name": tss[cid]["official_component_name"],
            "imo_parent_ref": tss[cid]["imo_parent_ref"],
            "imo_section": region,
            "legacy_ids": legacy_ids,
            "legacy_batch": "; ".join(sorted({r["Batch"] for r in recs})),
            "mapping_type": "; ".join(sorted({m["mapping_type"] for m in maps})) + ("; split from grouped legacy row" if split else ""),
            "proposed_association_status": status,
            "vts_name": svc_field(vts_ids, 1) or None,
            "vts_authority": svc_field(vts_ids, 3) or None,
            "vts_centre_sectors": svc_field(vts_ids, 4) or None,
            "vts_legacy_ids": "; ".join(vts_ids) or None,
            "vts_leads": svc_field(vts_lead_ids, 1) or None,
            "vts_boundary_basis": comp_note or " | ".join(r["VTS basis"] for r in recs if r["VTS basis"]) or None,
            "mandatory_reporting": "Yes" if mrs_ok else "Unresolved",
            "vrs_id": "; ".join(v for v, _ in imo_mrs) or None,
            "reporting_scheme_name": "; ".join([n for _, n in imo_mrs] + nat_mrs) or None,
            "reporting_boundary_basis": " | ".join(r["Reporting basis"] for r in recs if r["Reporting basis"]) or None,
            "voluntary_reporting": "Yes: " + "; ".join(sorted({schemes[x][1] for x in vol})) if vol else None,
            "applicability_notes": " | ".join(r["Applicability notes"] for r in recs if r["Applicability notes"]) or None,
            "legacy_source_urls": "\n".join(sources[s][3] for s in rec_src if s in sources and sources[s][3]),
            "open_questions": " | ".join(r["Open questions"] for r in recs if r["Open questions"]) or None,
            "carry_over_notes": " ".join(notes) or None,
            "imo2025_batch": f"B{(int(cid[4:]) - 1) // 5 + 1:02d}",
            "main_stage2_note": MAIN_B20.get(cid),
            "review_status": "Legacy lead: re-audit under Stage 2 before use",
            "carried_over": TODAY,
        })

    with open(OUT, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys()), lineterminator="\n")
        w.writeheader()
        w.writerows(rows)
    print(len(rows), "rows written")


if __name__ == "__main__":
    main()
